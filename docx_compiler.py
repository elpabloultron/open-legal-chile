"""Compilador de documentos Word (.docx) para los escritos y dossieres de la suite.

Regla de la casa: **los documentos se entregan en Word, no en PDF**, para que el abogado pueda
corregirlos. El PDF, cuando haga falta (presentación, notificación), se obtiene después desde Word
o con LibreOffice — pero el documento de trabajo es editable.

Convierte el Markdown del escrito (encabezados `#`, `##`, `###`, negritas `**…**`, viñetas)
en un .docx con tipografía forense: Times New Roman 12, cuerpo justificado.
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    _DOCX_DISPONIBLE = True
except ImportError:
    Document = None  # type: ignore[misc,assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[misc,assignment]
    Pt = None  # type: ignore[misc,assignment]
    _DOCX_DISPONIBLE = False


class WordDossierCompiler:
    """Convierte el Markdown de un escrito en un .docx editable."""

    def is_available(self) -> bool:
        """True si python-docx está instalado y se puede compilar Word."""
        return _DOCX_DISPONIBLE

    def compile(self, markdown_content: str, output_docx_path: str, title: str = "") -> Dict[str, Any]:
        """Escribe `markdown_content` como documento Word en `output_docx_path`."""
        if not _DOCX_DISPONIBLE or Document is None or Pt is None or WD_ALIGN_PARAGRAPH is None:
            return {"ok": False, "motivo": "python-docx no está instalado"}
        documento = Document()
        normal: Any = documento.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        if title:
            documento.core_properties.title = title
        documento.core_properties.author = "Open Legal Chile"
        for linea in markdown_content.split("\n"):
            self._agregar_linea(documento, linea)
        carpeta = os.path.dirname(os.path.abspath(output_docx_path))
        os.makedirs(carpeta, exist_ok=True)
        documento.save(output_docx_path)
        return {"ok": True, "ruta": output_docx_path, "caracteres": len(markdown_content)}

    # ── internos ────────────────────────────────────────────────────────────────────────────
    def _agregar_linea(self, documento: Any, linea: str) -> None:
        t = linea.rstrip()
        if not t.strip():
            return
        despojado = t.strip()
        if despojado.startswith("### "):
            documento.add_heading(despojado[4:].strip(), level=2)
        elif despojado.startswith("## "):
            documento.add_heading(despojado[3:].strip(), level=1)
        elif despojado.startswith("# "):
            documento.add_heading(despojado[2:].strip(), level=0)
        elif despojado in {"---", "***", "___"}:
            documento.add_page_break()
        elif despojado.startswith(("- ", "* ")) and len(despojado) > 2:
            parrafo = documento.add_paragraph(style="List Bullet")
            self._runs_con_formato(parrafo, despojado[2:].strip())
        else:
            parrafo = documento.add_paragraph()
            parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs_con_formato(parrafo, despojado)

    @staticmethod
    def _runs_con_formato(parrafo: Any, texto: str) -> None:
        """Negritas e itálicas del Markdown, como runs del párrafo."""
        for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
            if not parte:
                continue
            if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
                parrafo.add_run(parte[2:-2]).bold = True
            elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
                parrafo.add_run(parte[1:-1]).italic = True
            else:
                parrafo.add_run(parte)
